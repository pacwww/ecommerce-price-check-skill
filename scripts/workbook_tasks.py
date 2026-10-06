#!/usr/bin/env python3
"""Create row tasks from an ecommerce workbook and apply reviewed results."""
import argparse, json, math
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

HEADERS = ["ID", "商品名称", "规格", "跳转链接", "券后价", "划线价"]
EXTRA = ["店铺名称", "实际规格", "优惠条件", "查询时间", "状态", "备注", "本地截图路径"]

def load_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("需要 openpyxl：python3 -m pip install openpyxl") from exc
    return load_workbook(path)

def open_book(path):
    if path.suffix.lower() == ".xlsx": return load_xlsx(path), False
    if path.suffix.lower() != ".xls": raise SystemExit("仅支持 .xls 或 .xlsx")
    try:
        import xlrd
        from openpyxl import Workbook
    except ImportError as exc:
        raise SystemExit("读取 XLS 需要 xlrd 和 openpyxl") from exc
    old, new = xlrd.open_workbook(path), Workbook()
    new.remove(new.active)
    for sheet in old.sheets():
        ws = new.create_sheet(sheet.name)
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                value = sheet.cell_value(r, c)
                ws.cell(r + 1, c + 1).value = value if value != "" else None
    return new, True

def cell_url(cell):
    if getattr(cell, "hyperlink", None) and cell.hyperlink.target: return str(cell.hyperlink.target).strip()
    return str(cell.value or "").strip()

def tasks_from_book(book):
    tasks = []
    for ws in book.worksheets:
        actual = [str(ws.cell(1, i).value or "").strip() for i in range(1, 7)]
        if actual != HEADERS: raise SystemExit(f"工作表 {ws.title} 前六列应为 {HEADERS}，实际为 {actual}")
        for row in range(2, ws.max_row + 1):
            values = [ws.cell(row, col).value for col in range(1, 7)]
            url = cell_url(ws.cell(row, 4))
            if not any(v not in (None, "") for v in values[:4]) or not url: continue
            if urlparse(url).scheme not in {"http", "https"}: raise SystemExit(f"{ws.title} 第 {row} 行链接不是 HTTP(S)：{url}")
            tasks.append({"task_key":f"{ws.title}:{row}","sheet":ws.title,"row":row,"platform":ws.title,
                "id":values[0],"name":values[1],"spec":values[2],"url":url,
                "existing_coupon_price":values[4],"existing_original_price":values[5]})
    if len(tasks) > 100: raise SystemExit(f"本批共有 {len(tasks)} 行，超过 100 行限制")
    return tasks

def inspect(args):
    source = Path(args.input).expanduser().resolve(); book, legacy = open_book(source)
    tasks = tasks_from_book(book)
    payload = {"source":str(source),"legacy_xls":legacy,"generated_at":datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),"task_count":len(tasks),"tasks":tasks}
    output = Path(args.output).expanduser().resolve(); output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"生成 {len(tasks)} 个任务：{output}")

def valid_price(value, key):
    if value is None or value == "": return None
    if isinstance(value, bool): raise SystemExit(f"{key} 价格不能是布尔值")
    try: number = float(value)
    except (TypeError, ValueError): raise SystemExit(f"{key} 价格不是数字：{value!r}")
    if not math.isfinite(number) or number < 0: raise SystemExit(f"{key} 价格无效：{value!r}")
    return number

def apply(args):
    source = Path(args.input).expanduser().resolve(); output = Path(args.output).expanduser().resolve()
    if source == output: raise SystemExit("输出文件不能覆盖源文件")
    book, legacy = open_book(source); manifest = {t["task_key"]:t for t in tasks_from_book(book)}
    raw = json.loads(Path(args.results).read_text(encoding="utf-8")); results = raw.get("results") if isinstance(raw, dict) else raw
    if not isinstance(results, list): raise SystemExit("结果文件必须是 JSON 数组，或包含 results 数组")
    indexed = {}
    for item in results:
        key = item.get("task_key") if isinstance(item, dict) else None
        if key not in manifest: raise SystemExit(f"未知 task_key：{key!r}")
        if key in indexed: raise SystemExit(f"重复 task_key：{key}")
        indexed[key] = item
    fields = ["shop","actual_spec","conditions","queried_at","status","notes","screenshot_path"]
    for ws in book.worksheets:
        start = max(ws.max_column, 6) + 1
        for offset, title in enumerate(EXTRA): ws.cell(1, start + offset).value = title
        for key, task in manifest.items():
            if task["sheet"] != ws.title or key not in indexed: continue
            item, row = indexed[key], task["row"]
            if "coupon_price" in item: ws.cell(row, 5).value = valid_price(item.get("coupon_price"), key)
            if "original_price" in item: ws.cell(row, 6).value = valid_price(item.get("original_price"), key)
            for col in (5, 6): ws.cell(row, col).number_format = "0.00"
            for offset, field in enumerate(fields):
                if field in item: ws.cell(row, start + offset).value = str(item.get(field) or "")
    output.parent.mkdir(parents=True, exist_ok=True); book.save(output)
    if len(tasks_from_book(load_xlsx(output))) != len(manifest): raise SystemExit("导出复核失败：任务行数量发生变化")
    print(f"写入 {len(indexed)} 行结果：{output}" + ("（XLS 已转换为 XLSX）" if legacy else ""))

def main():
    parser = argparse.ArgumentParser(description=__doc__); sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("inspect"); p.add_argument("input"); p.add_argument("--output", required=True); p.set_defaults(func=inspect)
    p = sub.add_parser("apply"); p.add_argument("input"); p.add_argument("results"); p.add_argument("--output", required=True); p.set_defaults(func=apply)
    args = parser.parse_args(); args.func(args)

if __name__ == "__main__": main()
